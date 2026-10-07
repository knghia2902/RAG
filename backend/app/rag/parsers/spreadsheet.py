"""Spreadsheet parser for XLSX (streaming via openpyxl) and CSV (multiline RFC 4180)."""
import csv
import io
from typing import Any
from uuid import UUID
import openpyxl

from app.rag.contracts import EvidenceChunk, SourceLocator
from app.rag.parsers.base import (
    MalformedSourceError,
    ParsedRecord,
    ParseResult,
)


class SpreadsheetParser:
    """Parser adapter for XLSX and CSV files with row-level locators."""

    def parse_csv(self, content: bytes, locator_base: SourceLocator, dataset_id: UUID) -> ParseResult:
        try:
            text_content = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            try:
                text_content = content.decode("latin-1")
            except Exception as e:
                raise MalformedSourceError("Không thể giải mã tệp CSV với UTF-8 hoặc Latin-1") from e

        reader = csv.reader(io.StringIO(text_content))
        try:
            header = next(reader, None)
        except Exception as exc:
            raise MalformedSourceError(f"Lỗi đọc dữ liệu CSV: {exc}") from exc

        if header is None:
            return ParseResult(records=(), diagnostics=(), is_valid=True)

        header = [h.strip() if h is not None else f"col_{idx}" for idx, h in enumerate(header, start=1)]

        records: list[ParsedRecord] = []
        row_idx = 1  # header is row 1
        for row in reader:
            row_idx += 1
            if not row or not any(c.strip() for c in row if isinstance(c, str)):
                continue

            row_dict: dict[str, Any] = {}
            parts: list[str] = []
            for idx, h in enumerate(header):
                val = row[idx] if idx < len(row) else ""
                row_dict[h] = val
                parts.append(f"{h}: {val}")

            text = ", ".join(parts) if parts else "(dòng trống)"
            record_locator = SourceLocator(
                source_id=locator_base.source_id,
                source_version=locator_base.source_version,
                row=row_idx,
            )
            records.append(
                ParsedRecord(
                    position=row_idx,
                    locator=record_locator,
                    text=text,
                    metadata=row_dict,
                )
            )

        return ParseResult(records=tuple(records), diagnostics=(), is_valid=True)

    def parse_xlsx(self, content: bytes, locator_base: SourceLocator, dataset_id: UUID) -> ParseResult:
        records: list[ParsedRecord] = []
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise MalformedSourceError(f"Không thể mở tệp XLSX: {exc}") from exc

        try:
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, None)
                if not header:
                    continue

                col_names = [
                    str(h).strip() if h is not None else f"col_{i}"
                    for i, h in enumerate(header, start=1)
                ]

                row_idx = 1  # header is row 1
                for row_cells in rows:
                    row_idx += 1
                    if not row_cells or not any(c is not None for c in row_cells):
                        continue

                    row_dict: dict[str, Any] = {}
                    parts: list[str] = []
                    for idx, col_name in enumerate(col_names):
                        val = row_cells[idx] if idx < len(row_cells) else None
                        if val is not None:
                            val_str = str(val).strip()
                            row_dict[col_name] = val
                            parts.append(f"{col_name}: {val_str}")

                    if not parts:
                        continue

                    text = ", ".join(parts)
                    record_locator = SourceLocator(
                        source_id=locator_base.source_id,
                        source_version=locator_base.source_version,
                        sheet=sheet_name,
                        row=row_idx,
                    )
                    records.append(
                        ParsedRecord(
                            position=row_idx,
                            locator=record_locator,
                            text=text,
                            metadata=row_dict,
                        )
                    )
        finally:
            wb.close()

        return ParseResult(records=tuple(records), diagnostics=(), is_valid=True)

    def parse_bytes(self, content: bytes, locator_base: SourceLocator, dataset_id: UUID) -> ParseResult:
        # Detect ZIP header (0x50 0x4B 0x03 0x04) for XLSX
        if content.startswith(b"PK\x03\x04"):
            return self.parse_xlsx(content, locator_base, dataset_id)
        return self.parse_csv(content, locator_base, dataset_id)

    def parse(self, content: bytes, locator: SourceLocator, dataset_id: UUID) -> tuple[EvidenceChunk, ...]:
        res = self.parse_bytes(content, locator, dataset_id)
        return tuple(
            EvidenceChunk(dataset_id=dataset_id, locator=r.locator, text=r.text)
            for r in res.records
        )
