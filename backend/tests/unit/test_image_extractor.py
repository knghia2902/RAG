"""Unit tests for ImageExtractor service (PDF embedded/scan fallback & DOCX media extraction)."""
import io
import os
import shutil
import tempfile
from uuid import uuid4
import docx
from PIL import Image
import pypdfium2 as pdfium
import pytest

from app.rag.parsers.image_extractor import ImageExtractor, ExtractedImageData


@pytest.fixture
def temp_output_dir():
    dir_path = tempfile.mkdtemp(prefix="test_extracted_imgs_")
    yield dir_path
    shutil.rmtree(dir_path, ignore_errors=True)


PDF_WITH_TEXT_TEMPLATE = b"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>
endobj
4 0 obj
<< /Length 75 >>
stream
BT
/F1 12 Tf
100 700 Td
(Tai lieu huong dan xu ly su co TMS trang 1 he thong logistics WMS nang cao) Tj
ET
endstream
endobj
5 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000244 00000 n 
0000000370 00000 n 
trailer
<< /Size 6 /Root 1 0 R >>
startxref
447
%%EOF
"""


def create_sample_pdf_with_text_and_image(img_size=(100, 100)) -> bytes:
    """Tạo tệp PDF trong bộ nhớ chứa văn bản và một ảnh nhúng."""
    doc = pdfium.PdfDocument(io.BytesIO(PDF_WITH_TEXT_TEMPLATE))
    page = doc[0]

    # Thêm ảnh nhúng
    img_pil = Image.new("RGB", img_size, color="blue")
    pdf_img = pdfium.PdfImage.new(doc)
    pdf_img.set_bitmap(pdfium.PdfBitmap.from_pil(img_pil))
    pdf_img.set_matrix(pdfium.PdfMatrix(img_size[0], 0, 0, img_size[1], 50, 100))
    page.insert_obj(pdf_img)
    page.gen_content()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_sample_scanned_pdf_no_text() -> bytes:
    """Tạo tệp PDF mô phỏng tài liệu scan không có selectable text."""
    doc = pdfium.PdfDocument.new()
    page = doc.new_page(400, 500)
    # Trang chỉ vẽ hình ảnh mô phỏng tài liệu scan hoặc chụp màn hình
    img_pil = Image.new("RGB", (200, 250), color="yellow")
    pdf_img = pdfium.PdfImage.new(doc)
    pdf_img.set_bitmap(pdfium.PdfBitmap.from_pil(img_pil))
    pdf_img.set_matrix(pdfium.PdfMatrix(200, 0, 0, 250, 10, 10))
    page.insert_obj(pdf_img)
    page.gen_content()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_sample_docx_with_images(images: list[tuple[int, int, str]]) -> bytes:
    """Tạo tệp DOCX trong bộ nhớ với danh sách kích thước và màu sắc ảnh."""
    doc = docx.Document()
    doc.add_heading("Tài liệu SOP đính kèm hình ảnh", level=1)
    doc.add_paragraph("Các bước thực hiện thao tác trên hệ thống quản lý logistics:")

    for w, h, col in images:
        pil_img = Image.new("RGB", (w, h), color=col)
        img_buf = io.BytesIO()
        pil_img.save(img_buf, format="PNG")
        img_buf.seek(0)
        doc.add_picture(img_buf)

    out_buf = io.BytesIO()
    doc.save(out_buf)
    return out_buf.getvalue()


def test_extract_from_pdf_embedded_images(temp_output_dir):
    """Test trích xuất hình ảnh nhúng từ tệp PDF có text đầy đủ (> 50 ký tự)."""
    pdf_bytes = create_sample_pdf_with_text_and_image(img_size=(120, 120))

    dataset_id = uuid4()
    source_version_id = uuid4()
    extractor = ImageExtractor()

    images = extractor.extract_from_pdf(pdf_bytes, dataset_id, source_version_id, temp_output_dir)

    assert len(images) == 1
    img = images[0]
    assert isinstance(img, ExtractedImageData)
    assert img.dataset_id == dataset_id
    assert img.source_version_id == source_version_id
    assert img.page_number == 1
    assert img.image_type == "embedded"
    assert img.mime_type == "image/png"
    assert img.width == 120
    assert img.height == 120
    assert os.path.exists(img.storage_path)
    assert img.size_bytes > 0


def test_render_scanned_pdf_fallback_when_no_text(temp_output_dir):
    """Test fallback tự động render toàn trang scan thành bitmap chất lượng cao (D-02)."""
    pdf_bytes = create_sample_scanned_pdf_no_text()

    dataset_id = uuid4()
    source_version_id = uuid4()
    extractor = ImageExtractor()

    images = extractor.extract_from_pdf(pdf_bytes, dataset_id, source_version_id, temp_output_dir)

    assert len(images) == 1
    img = images[0]
    assert img.image_type == "rendered_page"
    assert img.page_number == 1
    assert os.path.exists(img.storage_path)
    # Vì scale=2.0 trên trang 400x500 nên kích thước bitmap kết quả là 800x1000
    assert img.width == 800
    assert img.height == 1000
    assert img.size_bytes > 0


def test_extract_media_from_docx(temp_output_dir):
    """Test trích xuất hình ảnh từ tệp DOCX (thư mục word/media/)."""
    docx_bytes = create_sample_docx_with_images([(150, 150, "green")])

    dataset_id = uuid4()
    source_version_id = uuid4()
    extractor = ImageExtractor()

    images = extractor.extract_from_docx(docx_bytes, dataset_id, source_version_id, temp_output_dir)

    assert len(images) == 1
    img = images[0]
    assert img.image_type == "embedded"
    assert img.page_number is None
    assert img.width == 150
    assert img.height == 150
    assert os.path.exists(img.storage_path)
    assert img.size_bytes > 0


def test_skip_small_icons_and_resize_oversized(temp_output_dir):
    """Test bỏ qua icon nhỏ (<64x64) và tự động co tỉ lệ ảnh quá lớn (>1600px)."""
    # 1 ảnh nhỏ 32x32 (sẽ bị bỏ qua) và 1 ảnh lớn 1800x900 (sẽ co về max 1600)
    docx_bytes = create_sample_docx_with_images([
        (32, 32, "gray"),
        (1800, 900, "purple"),
    ])

    dataset_id = uuid4()
    source_version_id = uuid4()
    extractor = ImageExtractor(max_dimension=1600, min_dimension=64)

    images = extractor.extract_from_docx(docx_bytes, dataset_id, source_version_id, temp_output_dir)

    # Chỉ có 1 ảnh được trích xuất (ảnh 32x32 bị bỏ qua)
    assert len(images) == 1
    img = images[0]
    assert img.width <= 1600
    assert img.height <= 1600
    assert img.width == 1600
    assert img.height == 800


def test_extractor_dispatcher_and_unsupported_formats(temp_output_dir):
    """Test hàm extract_images dispatcher tự nhận diện đuôi file."""
    dataset_id = uuid4()
    source_version_id = uuid4()
    extractor = ImageExtractor()

    # File không hỗ trợ ảnh (ví dụ txt, log) -> trả về danh sách rỗng
    txt_images = extractor.extract_images(b"hello world", "sample.txt", dataset_id, source_version_id, temp_output_dir)
    assert txt_images == []

    # File rỗng -> trả về rỗng
    empty_pdf = extractor.extract_images(b"", "empty.pdf", dataset_id, source_version_id, temp_output_dir)
    assert empty_pdf == []
